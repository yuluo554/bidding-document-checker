"""docx 评审报告生成（M5 交付，依赖 report extra；python-docx 惰性导入）。

报告结构（plan/04 §六 + HANDOFF-M5）：项目概况 → 结论汇总 → 分级问题明细
→ 依据与证据锚点（法规条文原文 + 招标条款原文，报告可溯源）→ 符合性矩阵
（符合项按需展开）→ 围标特征提示 → 签署栏 + 免责声明。

报告层口径（引擎只产非符合结论，D19；符合明细归报告层——HANDOFF-M5 DoD）：
- 「符合（未发现不符合）」= 规则门控激活 + 引擎数据前提成立 + 该 check_id 在该目标
  上无非符合结论；数据前提不成立（如报价卡/保证金响应缺失）→「未复核」（宁缺勿错
  同源，不虚计符合）；联动否决项（CHK-RJ-02/03/04）属条件触发，未触发不计符合；
- 输入：findings.json（run 产物）+ 同目录 cards/*.card.json + 项目原文
  （payload.project_dir 或 --project）；卡片/原文缺失 → 相应推导降级并在报告中注明
  （可解释降级），不阻断报告生成；
- **无外链纪律**：报告不写入任何 URL（法规来源台账见 data/knowledge/，报告只引用
  法规名+条号+条文原文）；内容不含时间戳（签署栏留空由评审人手填）。
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from bidchecker import __version__
from bidchecker.amounts import words_to_amount
from bidchecker.rules.engine import _CAP_RE
from bidchecker.schema import BidCard, TenderCard

_DEPENDENCY_HINT = "需要 report extra：pip install bidchecker[report]"
_REPORT_STAMP = f"bidchecker v{__version__}（辅助评审定位，非评标决策工具）"

# 侧别归属（镜像 engine.run_tender_checks / run_checks 的分派面；D19）
TENDER_SIDE_CHECKS: Tuple[str, ...] = ("CHK-DT-02",)
IMPLIED_REJECTION_CHECKS: Tuple[str, ...] = ("CHK-RJ-02", "CHK-RJ-03", "CHK-RJ-04")

# 符合性推导的规则类型列序（引擎 RULE_TYPES 同序）
MATRIX_RULE_TYPES: Tuple[str, ...] = (
    "qualification_match", "star_response", "arithmetic_consistency",
    "date_logic", "bond_check", "rejection_scan", "price_check",
)
_RULE_TYPE_ZH = {
    "qualification_match": "资格比对", "star_response": "星号条款",
    "arithmetic_consistency": "算术复核", "date_logic": "时间逻辑",
    "bond_check": "保证金", "rejection_scan": "否决扫描", "price_check": "价格提示",
}

# 未复核原因（数据前提不成立时的一句话说明；镜像引擎 handler 首行守卫）
_UNCHECKED_REASON = {
    "CHK-Q-01": "招标侧无资格/财务类要求行",
    "CHK-Q-02": "无带有效期的对应证明材料或要求截止",
    "CHK-Q-03": "无带等级要求的资格条件或投标侧未提供等级",
    "CHK-Q-04": "招标侧无业绩类要求行",
    "CHK-Q-05": "招标侧无人员类要求行",
    "CHK-SR-01": "招标侧无星号条款",
    "CHK-SR-02": "星号条款或响应值不可提取（宁缺勿错跳过）",
    "CHK-SR-03": "星号条款或响应值不可提取（宁缺勿错跳过）",
    "CHK-AR-01": "报价卡缺失或大小写金额不全",
    "CHK-AR-02": "报价卡缺失或分项合计/总价不全",
    "CHK-AR-03": "报价卡缺失或分项明细/分项合计不全",
    "CHK-BD-01": "缺保证金要求或缴纳响应数据",
    "CHK-BD-02": "缺保证金要求或缴纳形式数据",
    "CHK-DT-01": "有效期要求/承诺/否决条款数据不全",
    "CHK-DT-02": "开标时间或递交截止缺失",
    "CHK-RJ-01": "最高限价不可提取或报价缺失",
    "CHK-PC-01": "最高限价不可提取或报价缺失",
}

_ROW_CONFORM = "符合"
_ROW_UNCHECKED = "未复核"
_ROW_INACTIVE = "—"
_ROW_TRIGGER = "联动未触发"

_BASIS_LAW_RE = re.compile(r"^《(.+?)》(.+)$")
_BASIS_CLAUSE_RE = re.compile(r"^(Q\d+|S\d+|R\d+)$")


class ReportDependencyError(ImportError):
    """report extra 未安装（python-docx 缺失）。"""


@dataclass
class ReportResult:
    """报告生成结果摘要（CLI stdout 用）。"""

    path: str
    findings_total: int
    status_counts: Dict[str, int]
    conform_count: int = 0
    unchecked_count: int = 0
    degraded: List[str] = field(default_factory=list)


# ---------------------------------------------------------------- 符合性推导（报告层口径）
def _precondition(check_id: str, tender: TenderCard, bid: Optional[BidCard]) -> bool:
    """引擎数据前提镜像（engine.py 对应 handler 的守卫条件；报告层不虚计符合）。"""
    if check_id == "CHK-DT-02":
        return bool(tender.key_dates.open_time and tender.key_dates.submit_deadline)
    if bid is None:
        return False
    prices = bid.prices
    quals = tender.qualifications
    bid_quals = {q.item_id: q for q in (bid.qualifications or [])}
    if check_id == "CHK-Q-01":
        return any(q.category in ("qualification", "finance") for q in quals)
    if check_id == "CHK-Q-02":
        deadline = tender.key_dates.submit_deadline
        return any(q.category in ("qualification", "finance")
                   and (q.valid_until or deadline)
                   and q.item_id in bid_quals and bid_quals[q.item_id].valid_until
                   for q in quals)
    if check_id == "CHK-Q-03":
        return any(q.category in ("qualification", "finance") and q.level
                   and q.item_id in bid_quals for q in quals)
    if check_id == "CHK-Q-04":
        return any(q.category == "performance" for q in quals)
    if check_id == "CHK-Q-05":
        return any(q.category == "personnel" for q in quals)
    if check_id in ("CHK-SR-01",):
        return bool(tender.star_clauses)
    if check_id in ("CHK-SR-02", "CHK-SR-03"):
        from bidchecker.rules.engine import _star_value
        resp_by_id = {r.star_id: r for r in (bid.star_responses or [])}
        return any(_star_value(c.text) is not None
                   and c.item_id in resp_by_id
                   and _star_value(resp_by_id[c.item_id].quote or "") is not None
                   for c in tender.star_clauses)
    if check_id == "CHK-AR-01":
        if prices is None or not prices.total_words or prices.total_number is None:
            return False
        try:
            words_to_amount(prices.total_words)
        except ValueError:
            return False
        return True
    if check_id == "CHK-AR-02":
        return prices is not None and prices.items_sum is not None \
            and prices.total_number is not None
    if check_id == "CHK-AR-03":
        return prices is not None and bool(prices.items) and prices.items_sum is not None
    if check_id == "CHK-BD-01":
        return tender.bond is not None and bid.bond_response is not None \
            and tender.bond.amount is not None and bid.bond_response.amount is not None
    if check_id == "CHK-BD-02":
        return tender.bond is not None and bid.bond_response is not None \
            and bool(bid.bond_response.forms)
    if check_id == "CHK-DT-01":
        return bool(tender.key_dates.validity_days and bid.key_dates.validity_days is not None
                    and any("有效期" in r.text for r in tender.rejection_clauses))
    if check_id in ("CHK-RJ-01", "CHK-PC-01"):
        return _cap_extractable(tender) and prices is not None \
            and prices.total_number is not None
    return False


def _cap_extractable(tender: TenderCard) -> bool:
    """最高限价可提取（engine._price_cap 同口径：R 条款内嵌人民币金额）。"""
    for r in tender.rejection_clauses:
        if "最高限价" in r.text and _CAP_RE.search(r.text):
            return True
    return False


def _gate_active(rule: Optional[Dict[str, Any]], gate_text: str) -> bool:
    """only_if 门控（与引擎 rule_applies 同源）；门控文本不可用时对带 only_if 规则判未激活。"""
    from bidchecker.rules.gates import rule_applies
    if rule is None:
        return True
    return rule_applies(rule, gate_text)


def derive_conformance(payload: Dict[str, Any], rules: List[Dict[str, Any]],
                       tender_card: Optional[TenderCard],
                       bid_cards: Dict[str, BidCard],
                       texts: Dict[str, str]) -> List[Dict[str, Any]]:
    """全部 (目标 × check_id) 推导行：结论状态 + 报告层符合性（矩阵与符合项明细共用）。

    行字段：target/side/check_id/name/rule_type/status/conform/unchecked/reason/findings。
    status ∈ 符合(引擎结论语义仅非符合，此处为报告层推导)/不符合/待人工确认；
    conform ∈ _ROW_*。
    """
    by_id = {r.get("id"): r for r in rules if isinstance(r, dict) and r.get("id")}
    findings = payload.get("findings", [])
    tender_text = texts.get("tender.md")
    targets: List[str] = []
    for f in findings:
        t = f.get("target", "")
        if "|" not in t and t not in targets:
            targets.append(t)
    if tender_card is not None and "tender.md" not in targets:
        targets.insert(0, "tender.md")
    for t in sorted(bid_cards):
        if t not in targets:
            targets.append(t)
    targets.sort(key=lambda t: (t != "tender.md", t))

    rows: List[Dict[str, Any]] = []
    for target in targets:
        is_tender = target == "tender.md"
        side_checks = TENDER_SIDE_CHECKS if is_tender else _bid_side_checks()
        card: Any = tender_card if is_tender else bid_cards.get(target)
        if is_tender:
            gate_text = tender_text
        else:
            tt = tender_text or ""
            gate_text = (tt + "\n" + texts[target]) if target in texts else (tt or None)
        for check_id in side_checks:
            rule = by_id.get(check_id)
            fds = [f for f in findings
                   if f.get("target") == target and f.get("check_id") == check_id]
            name = (rule or {}).get("name", "")
            rule_type = (rule or {}).get("rule_type", "")
            row = {"target": target, "side": "tender" if is_tender else "bid",
                   "check_id": check_id, "name": name, "rule_type": rule_type,
                   "findings": fds, "status": "", "conform": _ROW_INACTIVE,
                   "unchecked": False, "reason": ""}
            if fds:
                if any(f.get("status") == "不符合" for f in fds):
                    row["status"] = "不符合"
                else:
                    row["status"] = "待人工确认"
                row["conform"] = "不符合" if row["status"] == "不符合" else "待人工确认"
            elif tender_card is None:
                # 招标卡缺失：数据前提无从判定（报告层宁缺勿错，不虚计符合）
                row["conform"] = _ROW_UNCHECKED
                row["unchecked"] = True
                row["reason"] = "招标卡缺失，符合性推导降级"
            elif check_id in IMPLIED_REJECTION_CHECKS:
                row["conform"] = _ROW_TRIGGER
            elif not _gate_active(rule, gate_text or ""):
                row["conform"] = _ROW_INACTIVE
            elif card is not None and _precondition(check_id, tender_card, card):
                row["status"] = _ROW_CONFORM
                row["conform"] = _ROW_CONFORM
            else:
                row["conform"] = _ROW_UNCHECKED
                row["unchecked"] = True
                row["reason"] = _UNCHECKED_REASON.get(check_id, "引擎数据前提不成立")
            rows.append(row)
    return rows


def _bid_side_checks() -> Tuple[str, ...]:
    from bidchecker.rules.engine import CHECK_IDS, CHK_CL_PAIRS, CHK_DT_ORDER
    return tuple(cid for cid in CHECK_IDS
                 if cid not in TENDER_SIDE_CHECKS and cid != CHK_CL_PAIRS
                 and cid != CHK_DT_ORDER)


# ---------------------------------------------------------------- 报告输入装载
def load_report_inputs(findings_path, project_dir: Optional[str] = None
                       ) -> Tuple[Dict[str, Any], Optional[TenderCard],
                                  Dict[str, BidCard], Dict[str, str], List[str]]:
    """findings.json → (payload, tender卡, {目标: 投标卡}, {目标: 原文}, 降级说明)。

    卡片缺失/原文目录不可用不抛错（返回部分输入 + 降级说明，报告照常生成）。
    """
    path = Path(findings_path)
    if not path.is_file():
        raise FileNotFoundError(f"findings.json 不存在：{path}")
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except ValueError as exc:
        raise ValueError(f"findings.json 不是合法 JSON：{exc}") from None
    degraded: List[str] = []
    cards_dir = path.parent / "cards"
    tender_card: Optional[TenderCard] = None
    bid_cards: Dict[str, BidCard] = {}
    if (cards_dir / "tender.card.json").is_file():
        tender_card = TenderCard.from_dict(
            json.loads((cards_dir / "tender.card.json").read_text(encoding="utf-8")))
    else:
        degraded.append("卡片目录缺失（cards/tender.card.json），依据锚点与符合性推导降级")
    for card_path in sorted(cards_dir.glob("*.card.json")):
        if card_path.name == "tender.card.json":
            continue
        stem = card_path.name[: -len(".card.json")]
        bid_cards[f"bids/{stem}.md"] = BidCard.from_dict(
            json.loads(card_path.read_text(encoding="utf-8")))
    root = Path(project_dir) if project_dir else Path(str(payload.get("project_dir", "")))
    texts: Dict[str, str] = {}
    if root.is_dir():
        tp = root / "tender.md"
        if tp.is_file():
            texts["tender.md"] = tp.read_text(encoding="utf-8")
        bids_dir = root / "bids"
        if bids_dir.is_dir():
            for bp in sorted(bids_dir.glob("*.md")):
                texts[f"bids/{bp.name}"] = bp.read_text(encoding="utf-8")
    else:
        degraded.append("项目原文目录不可用，类目门控与原文锚点降级（结论明细不受影响）")
    return payload, tender_card, bid_cards, texts, degraded


# ---------------------------------------------------------------- docx 渲染
def _qn(tag: str) -> str:
    from docx.oxml.ns import qn
    return qn(tag)


def _set_cjk(style, east_asia: str = "SimSun") -> None:
    """样式补充中文字体（默认模板西文 Calibri 对中文标题渲染不稳）。"""
    style.font.name = "Calibri"
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(_qn("w:eastAsia"), east_asia)


def render_report(payload: Dict[str, Any], tender_card: Optional[TenderCard],
                  bid_cards: Dict[str, BidCard], texts: Dict[str, str],
                  out_path, *, include_pass: bool = False,
                  degraded: Optional[List[str]] = None) -> ReportResult:
    """payload + 卡片 + 原文 → docx 报告（无外链、无时间戳内容）。"""
    try:
        from docx import Document
    except ImportError as exc:
        raise ReportDependencyError(_DEPENDENCY_HINT) from exc
    degraded = list(degraded or [])
    from docx.shared import Pt

    doc = Document()
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        try:
            _set_cjk(doc.styles[style_name])
        except KeyError:
            continue
    doc.styles["Normal"].font.size = Pt(10.5)
    core = doc.core_properties
    core.author = "bidchecker"
    core.title = "招投标文件智能评审报告"
    core.comments = "辅助评审定位，非评标决策工具"

    rules = _load_rules_safe()
    by_id = {r.get("id"): r for r in rules if isinstance(r, dict) and r.get("id")}
    findings: List[Dict[str, Any]] = payload.get("findings", [])
    summary: Dict[str, Any] = payload.get("summary", {})
    counts: Dict[str, int] = summary.get("status_counts", {})

    # ---- 标题
    doc.add_heading("招投标文件智能评审报告", level=0)
    project = tender_card.project if tender_card is not None else None
    sub = project.name if (project is not None and project.name) \
        else str(payload.get("project_dir", ""))
    p = doc.add_paragraph(f"{sub} ｜ {_REPORT_STAMP}")
    p.runs[0].font.size = Pt(9)

    # ---- 一、项目概况
    doc.add_heading("一、项目概况", level=1)
    if project is not None:
        for label, value in (("项目名称", project.name), ("项目编号", project.number),
                             ("招标人（采购人）", project.buyer), ("代理机构", project.agency)):
            if value:
                doc.add_paragraph(f"{label}：{value}", style="List Bullet")
    n_bids = sum(1 for t in bid_cards if t.startswith("bids/"))
    scope = [f"评审范围：招标文件 1 份、投标文件 {n_bids} 份"]
    for t in sorted(bid_cards):
        scope.append(f"  · {bid_cards[t].bidder or t}（{t}）")
    doc.add_paragraph("\n".join(scope))
    failed = payload.get("summary", {}).get("failed") or []
    if failed:
        doc.add_paragraph(f"解析失败清单（结论覆盖不含这些文档）：{'、'.join(failed)}")
    if degraded:
        for d in degraded:
            doc.add_paragraph(f"降级说明：{d}")
    risk = {"high": 0, "mid": 0, "low": 0}
    for f in findings:
        if f.get("status") == "不符合" and f.get("risk_level") in risk:
            risk[f["risk_level"]] += 1
    doc.add_paragraph(
        f"评审结论 {len(findings)} 条：不符合 {counts.get('不符合', 0)}"
        f"（高风险 {risk['high']} / 中风险 {risk['mid']} / 低风险 {risk['low']}）、"
        f"待人工确认 {counts.get('待人工确认', 0)}。")

    # ---- 二、结论汇总
    doc.add_heading("二、结论汇总", level=1)
    t = doc.add_table(rows=4, cols=3)
    t.style = "Table Grid"
    for i, (h0, h1, h2) in enumerate((("判定", "数量", "说明"),)):
        pass
    _fill_row(t.rows[0], "判定", "数量", "说明")
    _fill_row(t.rows[1], "不符合", str(counts.get("不符合", 0)),
              "确定性规则检出，风险分级见问题明细")
    _fill_row(t.rows[2], "待人工确认", str(counts.get("待人工确认", 0)),
              "文本语义类提示或嫌疑线索，提请人工核查")
    _fill_row(t.rows[3], "符合", "（见符合性矩阵）", "引擎只产非符合结论；符合明细由报告层推导")

    # ---- 三、问题明细（分级）
    doc.add_heading("三、问题明细（分级）", level=1)
    rank = {"high": 0, "mid": 1, "low": 2}
    problems = [f for f in findings if f.get("status") in ("不符合", "待人工确认")]
    problems.sort(key=lambda f: (rank.get(f.get("risk_level"), 3),
                                 f.get("status") != "不符合",
                                 f.get("target", ""), f.get("check_id", "")))
    if not problems:
        doc.add_paragraph("未检出不符合项与待人工确认项（干净对照或全部复核通过）。")
    else:
        section_titles = {(0, True): "3.1 不符合 —— 高风险", (1, True): "3.2 不符合 —— 中风险",
                          (2, True): "3.3 不符合 —— 低风险", (0, False): "3.4 待人工确认"}
        seen = set()
        for f in problems:
            is_nc = f.get("status") == "不符合"
            key = (rank.get(f.get("risk_level"), 3), is_nc)
            if key not in seen:
                seen.add(key)
                doc.add_heading(section_titles.get(key, "3.4 待人工确认"), level=2)
            _render_finding(doc, f, by_id)

    # ---- 四、依据与证据锚点
    doc.add_heading("四、依据与证据锚点", level=1)
    doc.add_paragraph(
        "结论依据分两类：法规条文（源自知识库已核对条文块，官方逐字核验；来源台账见"
        " data/knowledge/raw/，报告不内嵌外链）与招标文件条款（按条款 id 锚定原文）。")
    _render_basis_anchors(doc, findings, tender_card)

    # ---- 五、符合性矩阵（报告层推导；符合项按需展开）
    doc.add_heading("五、符合性校核矩阵（报告层推导）", level=1)
    rows = derive_conformance(payload, rules, tender_card, bid_cards, texts)
    _render_matrix(doc, rows)
    doc.add_paragraph(
        "口径：「符合」= 规则门控激活且引擎数据前提成立且未产生不符合结论（即未发现不符合）；"
        "「未复核」= 数据前提不成立（宁缺勿错，不虚计符合）；「—」= 规则未激活；"
        "联动否决项（CHK-RJ-02/03/04）属条件触发，未触发不计入符合；围标特征为按对结论，"
        "不在矩阵逐格展示（见围标章节）。")
    conform_rows = [r for r in rows if r["conform"] == _ROW_CONFORM]
    unchecked_rows = [r for r in rows if r["conform"] == _ROW_UNCHECKED]
    if include_pass:
        doc.add_heading("5.1 符合项明细（--all 展开）", level=2)
        if conform_rows:
            for r in conform_rows:
                doc.add_paragraph(
                    f"{r['target']} ｜ {r['check_id']} {r['name']}：符合（未发现不符合）",
                    style="List Bullet")
        else:
            doc.add_paragraph("（无）")
        doc.add_heading("5.2 未能复核项", level=2)
        if unchecked_rows:
            for r in unchecked_rows:
                doc.add_paragraph(
                    f"{r['target']} ｜ {r['check_id']} {r['name']}：未复核（{r['reason']}）",
                    style="List Bullet")
        else:
            doc.add_paragraph("（无）")

    # ---- 六、围标特征提示
    doc.add_heading("六、围标特征提示", level=1)
    collusion = [f for f in findings if f.get("check_id") == "CHK-CL-01"]
    if collusion:
        for f in collusion:
            doc.add_paragraph(f"围标对 {f.get('target', '')}", style="List Bullet")
            doc.add_paragraph(f.get("message", ""))
    else:
        doc.add_paragraph("未检出异常一致文本特征（同分钟编制时间戳 + 对照独有共享长句，"
                          "双信号交叉验证；投标不足 3 份亦不产出）。")

    # ---- 签署栏 + 免责声明
    doc.add_heading("签署栏", level=1)
    doc.add_paragraph("评审人（签字）：________________    日期：________________")
    doc.add_heading("免责声明", level=1)
    doc.add_paragraph(
        "本报告由 bidchecker 自动生成，输出仅为辅助评审定位，不构成评标结论、法律意见或"
        "授标建议；数值类结论来自确定性规则并可溯源至法规/招标文件条款，围标分析仅限文本"
        "特征、不出定性结论；最终评审与决策应由依法组建的评标委员会作出。")

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return ReportResult(path=str(out), findings_total=len(findings),
                        status_counts=dict(counts), conform_count=len(conform_rows),
                        unchecked_count=len(unchecked_rows), degraded=degraded)


def _render_finding(doc, f: Dict[str, Any], by_id: Dict[str, Dict[str, Any]]) -> None:
    rule = by_id.get(f.get("check_id", ""))
    risk = f.get("risk_level")
    head = f"【{f.get('check_id', '')}】{(rule or {}).get('name', '')}"
    if risk:
        head += f"（风险 {risk}）"
    head += f" —— 目标 {f.get('target', '')}"
    p = doc.add_paragraph()
    run = p.add_run(head)
    run.bold = True
    doc.add_paragraph(f"问题：{f.get('message', '')}")
    basis = f.get("basis") or []
    if basis:
        doc.add_paragraph("依据：" + "；".join(basis))
    if rule and rule.get("advice"):
        doc.add_paragraph(f"建议：{rule['advice']}")
    for ev in f.get("evidence") or []:
        loc = ev.get("section") or ""
        line = f"（第 {ev['line']} 行）" if ev.get("line") else ""
        quote = ev.get("quote") or ""
        doc.add_paragraph(f"证据：{loc}{line}：“{quote}”")
    if not (f.get("evidence") or []):
        doc.add_paragraph("证据：——（数值/逻辑规则直接判定）")


def _render_basis_anchors(doc, findings: List[Dict[str, Any]],
                          tender_card: Optional[TenderCard]) -> None:
    law_refs: List[Tuple[str, str]] = []
    clause_ids: List[str] = []
    for f in findings:
        for b in f.get("basis") or []:
            m = _BASIS_LAW_RE.match(b)
            if m:
                ref = (m.group(1), m.group(2))
                if ref not in law_refs:
                    law_refs.append(ref)
            elif _BASIS_CLAUSE_RE.match(b) and b not in clause_ids:
                clause_ids.append(b)
    if law_refs:
        doc.add_heading("4.1 法规条文", level=2)
        blocks = {(str(b.get("law", "")), str(b.get("article_no", ""))): b
                  for b in _load_blocks_safe()}
        for law, article in law_refs:
            block = blocks.get((law, article))
            if block is None:
                doc.add_paragraph(f"《{law}》{article}：知识库未收录该条文块（依据关联率"
                                  "由引擎测试把关；本行为兜底渲染）", style="List Bullet")
                continue
            doc.add_paragraph(f"《{law}》{article}【{block.get('status', '')}】",
                              style="List Bullet")
            doc.add_paragraph(str(block.get("text", "")))
    if clause_ids:
        doc.add_heading("4.2 招标文件条款", level=2)
        index: Dict[str, Tuple[str, Any]] = {}
        if tender_card is not None:
            for q in tender_card.qualifications:
                index[q.item_id] = ("资格要求", q)
            for s in tender_card.star_clauses:
                index[s.item_id] = ("星号条款", s)
            for r in tender_card.rejection_clauses:
                index[r.item_id] = ("否决条款", r)
        for cid in clause_ids:
            hit = index.get(cid)
            if hit is None:
                doc.add_paragraph(f"{cid}：招标卡缺失或条款 id 未定位", style="List Bullet")
                continue
            kind, clause = hit
            text = clause.requirement if kind == "资格要求" else clause.text
            line = f"（原文第 {clause.evidence.line} 行）" \
                if clause.evidence is not None and clause.evidence.line else ""
            doc.add_paragraph(f"{cid}（{kind}）{line}：{text}", style="List Bullet")
    if not law_refs and not clause_ids:
        doc.add_paragraph("（结论无依据引用）")


def _render_matrix(doc, rows: List[Dict[str, Any]]) -> None:
    targets = sorted({r["target"] for r in rows}, key=lambda t: (t != "tender.md", t))
    head = ["目标 \\ 规则类型"] + [_RULE_TYPE_ZH.get(rt, rt) for rt in MATRIX_RULE_TYPES]
    table = doc.add_table(rows=len(targets) + 1, cols=len(head))
    table.style = "Table Grid"
    for c, htext in enumerate(head):
        table.rows[0].cells[c].text = htext
    for ri, target in enumerate(targets, start=1):
        table.rows[ri].cells[0].text = target
        for ci, rt in enumerate(MATRIX_RULE_TYPES, start=1):
            cell_rows = [r for r in rows if r["target"] == target and r["rule_type"] == rt]
            table.rows[ri].cells[ci].text = _matrix_cell(cell_rows)
    doc.add_paragraph("围标特征提示为按对结论（CHK-CL-01），不进入矩阵逐格。")


def _matrix_cell(cell_rows: List[Dict[str, Any]]) -> str:
    if not cell_rows:
        return _ROW_INACTIVE
    n_nc = sum(1 for r in cell_rows if r["conform"] == "不符合")
    n_manual = sum(1 for r in cell_rows if r["conform"] == "待人工确认")
    if n_nc:
        return f"不符合 {n_nc}"
    if n_manual:
        return f"待人工 {n_manual}"
    if any(r["conform"] == _ROW_CONFORM for r in cell_rows):
        return _ROW_CONFORM
    if any(r["conform"] == _ROW_UNCHECKED for r in cell_rows):
        return _ROW_UNCHECKED
    return _ROW_TRIGGER if all(r["conform"] == _ROW_TRIGGER for r in cell_rows) else _ROW_INACTIVE


def _fill_row(row, *values: str) -> None:
    for cell, value in zip(row.cells, values):
        cell.text = value


def _load_rules_safe() -> List[Dict[str, Any]]:
    try:
        from bidchecker.knowledge import load_rules
        return list(load_rules())
    except Exception:  # noqa: BLE001 —— 知识库缺失时报告仍可生成（规则名/建议降级为空）
        return []


def _load_blocks_safe() -> List[Dict[str, Any]]:
    try:
        from bidchecker.knowledge import load_blocks
        return list(load_blocks())
    except Exception:  # noqa: BLE001 —— 同上，锚点降级为条款 id 直引
        return []


# ---------------------------------------------------------------- 入口
def build_report(findings_path, out_path=None, *, project_dir: Optional[str] = None,
                 include_pass: bool = False) -> ReportResult:
    """findings.json（run 产物）→ docx 评审报告。

    out_path 缺省 = findings.json 同目录 report.docx；卡片与项目原文自动就近装载，
    缺失部分降级并在报告中注明。docx 内容无外链、无时间戳。
    """
    payload, tender_card, bid_cards, texts, degraded = load_report_inputs(
        findings_path, project_dir=project_dir)
    out = Path(out_path) if out_path else Path(findings_path).parent / "report.docx"
    return render_report(payload, tender_card, bid_cards, texts, out,
                         include_pass=include_pass, degraded=degraded)


__all__ = ["ReportDependencyError", "ReportResult", "build_report", "load_report_inputs",
           "render_report", "derive_conformance"]
