"""docx 评审报告（M5）：结构、docx 无外链断言、依据锚点、符合性推导口径、降级路径。"""
import re
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("docx")

from bidchecker.knowledge import load_blocks, load_rules
from bidchecker.pipeline import run_pipeline
from bidchecker.report import (ReportDependencyError, build_report,
                               derive_conformance, load_report_inputs)

ROOT = Path(__file__).resolve().parent.parent
EXAMPLES = ROOT / "examples"


@pytest.fixture(scope="module")
def run_out(tmp_path_factory):
    """demo-goods-1 的 run 产物（findings.json + cards/），模块级复用。"""
    out = tmp_path_factory.mktemp("run") / "output"
    run_pipeline(EXAMPLES / "demo-goods-1", out_dir=str(out))
    return out


def _docx_text(path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


# ---------------------------------------------------------------- 基本结构与内容
def test_build_report_structure(run_out, tmp_path):
    rr = build_report(run_out / "findings.json", tmp_path / "report.docx",
                      include_pass=True)
    assert Path(rr.path).is_file()
    assert rr.findings_total == 8
    assert rr.status_counts["不符合"] == 7 and rr.status_counts["待人工确认"] == 1
    assert rr.conform_count > 0
    full = _docx_text(rr.path)
    for head in ("招投标文件智能评审报告", "一、项目概况", "二、结论汇总",
                 "三、问题明细（分级）", "四、依据与证据锚点", "五、符合性校核矩阵",
                 "六、围标特征提示", "签署栏", "免责声明",
                 "5.1 符合项明细", "5.2 未能复核项"):
        assert head in full, f"报告缺章节：{head}"
    # 问题明细分级标题
    assert "3.1 不符合 —— 高风险" in full and "3.2 不符合 —— 中风险" in full
    assert "3.4 待人工确认" in full
    # 结论内容与依据进入报告
    assert "CHK-DT-02" in full and "CHK-CL-01" in full
    assert "围标对 bids/启明智能化工程有限公司.md|bids/安捷计算机系统有限公司.md" in full


def test_docx_zero_external_links(run_out, tmp_path):
    """docx 无外链硬断言：无外部 relationship、无 hyperlink 元素、正文无 http 痕迹。"""
    rr = build_report(run_out / "findings.json", tmp_path / "r.docx")
    with zipfile.ZipFile(rr.path) as z:
        for name in z.namelist():
            if name.endswith(".rels"):
                content = z.read(name).decode("utf-8")
                assert 'TargetMode="External"' not in content, f"{name} 含外链关系"
        document = z.read("word/document.xml").decode("utf-8")
    assert "<w:hyperlink" not in document
    text = re.sub(r"<[^>]+>", "", document)
    assert "http" not in text, "报告正文不得包含 URL（来源台账见 data/knowledge/）"
    assert "www." not in text


def test_basis_anchors_resolve_to_blocks_and_clauses(run_out, tmp_path):
    """依据关联链：Finding.basis → 条文块原文 / 招标条款原文，全部落进报告。"""
    payload, _tender, _bids, _texts, _deg = load_report_inputs(run_out / "findings.json")
    blocks = {(str(b.get("law", "")), str(b.get("article_no", ""))): b
              for b in load_blocks()}
    rr = build_report(run_out / "findings.json", tmp_path / "r.docx")
    full = _docx_text(rr.path)
    checked = 0
    for f in payload["findings"]:
        for b in f.get("basis") or []:
            m = re.match(r"^《(.+?)》(.+)$", b)
            if m:
                block = blocks.get((m.group(1), m.group(2)))
                assert block is not None, f"知识库缺条文块：{b}"
                assert block["text"] in full, f"条文原文未进报告：{b}"
                checked += 1
            elif re.match(r"^(Q\d+|S\d+|R\d+)$", b):
                assert b in full, f"条款 id 未进报告：{b}"
                checked += 1
    assert checked > 0


def test_conformance_rows_consistent_with_findings(run_out):
    """报告层推导口径：有结论的 (目标, check_id) 行 = 结论状态；其余按门控/前提推导。"""
    payload, tender, bids, texts, _deg = load_report_inputs(run_out / "findings.json")
    rows = derive_conformance(payload, load_rules(), tender, bids, texts)
    by_key = {}
    for f in payload["findings"]:
        if "|" in f["target"]:
            continue  # 围标对不进矩阵
        by_key.setdefault((f["target"], f["check_id"]), []).append(f)
    for row in rows:
        fds = by_key.get((row["target"], row["check_id"]))
        if fds:
            expect = "不符合" if any(f["status"] == "不符合" for f in fds) else "待人工确认"
            assert row["status"] == expect and row["conform"] == expect
        else:
            assert row["status"] in ("", "符合")
    # 推导存在符合项；联动否决项不记符合
    assert any(r["conform"] == "符合" for r in rows)
    for row in rows:
        if row["check_id"] in ("CHK-RJ-02", "CHK-RJ-03", "CHK-RJ-04"):
            assert row["conform"] != "符合"


def test_clean_project_matrix_all_conform(tmp_path):
    out = tmp_path / "out"
    run_pipeline(EXAMPLES / "demo-goods-2", out_dir=str(out))
    payload, _tender, _bids, _texts, _deg = load_report_inputs(out / "findings.json")
    rows = derive_conformance(payload, load_rules(), _tender, _bids, _texts)
    assert payload["findings"] == []
    conforms = [r for r in rows if r["conform"] == "符合"]
    assert conforms, "干净项目应推导出符合项"
    rr = build_report(out / "findings.json", tmp_path / "clean.docx")
    assert rr.findings_total == 0 and rr.conform_count == len(conforms)
    full = _docx_text(rr.path)
    assert "未检出不符合项与待人工确认项" in full


# ---------------------------------------------------------------- 降级路径
def test_degraded_without_cards(tmp_path, run_out):
    """只有 findings.json（无 cards/）也能生成报告：锚点与符合性推导降级并注明。"""
    lone = tmp_path / "lone"
    lone.mkdir()
    (lone / "findings.json").write_bytes((run_out / "findings.json").read_bytes())
    rr = build_report(lone / "findings.json", lone / "r.docx")
    assert rr.findings_total == 8
    assert rr.degraded and any("卡片目录缺失" in d for d in rr.degraded)
    full = _docx_text(rr.path)
    assert "降级说明" in full
    # 无卡片 → 符合性推导降级为未复核（不虚计符合）
    assert rr.conform_count == 0


def test_degraded_without_project_texts(tmp_path, run_out):
    import json
    import shutil

    lone = tmp_path / "lone2"
    lone.mkdir()
    payload = json.loads((run_out / "findings.json").read_text(encoding="utf-8"))
    payload["project_dir"] = str(tmp_path / "no-such-project")
    (lone / "findings.json").write_bytes(
        (json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
    shutil.copytree(run_out / "cards", lone / "cards")
    rr = build_report(lone / "findings.json", lone / "r.docx")
    assert any("项目原文目录不可用" in d for d in rr.degraded)
    # 卡片在 → 符合性推导可用（当前 rules.json 均无 only_if）
    assert rr.conform_count > 0


def test_missing_findings_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        build_report(tmp_path / "no-such.json", tmp_path / "r.docx")


def test_invalid_findings_json_raises(tmp_path, run_out):
    p = tmp_path / "bad.json"
    p.write_bytes(b"{not-json")
    with pytest.raises(ValueError):
        build_report(p, tmp_path / "r.docx")


def test_report_dependency_error_hint(tmp_path, run_out, monkeypatch):
    """report extra 未安装 → ReportDependencyError（带 pip 安装提示），CLI 退出码 2 消费。"""
    import sys

    monkeypatch.setitem(sys.modules, "docx", None)
    with pytest.raises(ReportDependencyError) as exc:
        build_report(run_out / "findings.json", tmp_path / "r.docx")
    assert "pip install bidchecker[report]" in str(exc.value)


def test_project_dir_override(tmp_path, run_out):
    """--project 覆盖 findings.json 内 project_dir（findings 挪位后仍可完整推导）。"""
    import shutil

    moved = tmp_path / "moved"
    moved.mkdir()
    shutil.copytree(run_out / "cards", moved / "cards")
    shutil.copy(run_out / "findings.json", moved / "findings.json")
    rr = build_report(moved / "findings.json", moved / "r.docx",
                      project_dir=str(EXAMPLES / "demo-goods-1"))
    assert rr.conform_count > 0
    assert not any("项目原文目录不可用" in d for d in rr.degraded)


def test_report_no_timestamp_content(run_out, tmp_path):
    """报告内容无时间戳（签署栏留空手填；日期仅来自证据原文）。"""
    rr = build_report(run_out / "findings.json", tmp_path / "r.docx")
    full = _docx_text(rr.path)
    assert "生成时间" not in full and "生成日期" not in full
    assert "评审人（签字）：" in full and "日期：____" in full
